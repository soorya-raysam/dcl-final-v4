#!/usr/bin/env python
# coding: utf-8

# ## FINAL GUI

# In[1]:


import os
import re
import time
import string
import paramiko
import openpyxl
from datetime import datetime
from openpyxl.styles import Font, Alignment
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.action_chains import ActionChains

# Config
CHROME_DRIVER_PATH = "/Users/sooryaraysam/Downloads/chromedriver-mac-arm64/chromedriver"
LOGIN_URL = "https://172.16.70.39/login/webLogin"
WEB_USERNAME = "dadmin"
WEB_PASSWORD = "Avaya@123"
SAT_HOST = "172.16.70.39"
SAT_PORT = 5022
SAT_USERNAME = "dadmin"
SAT_PASSWORD = "Avaya@123"
SSH_HOST = '172.16.70.20'
SSH_PORT = 22
SSH_USERNAME = 'cust'
SSH_PASSWORD = 'Avaya@123'
excel_path = "avaya_page_log.xlsx"

web_pages = {
    "Current Alarms": "/cgi-bin/cm/alrmCurrentAlarms/w_currentAlarms",
    "Status Summary": "/cgi-bin/cm/servStatusSummary/w_statusSummary",
    "Server Date/Time": "/cgi-bin/cm/servDateTime/w_serverDateTime",
    "Software Version": "/cgi-bin/cm/servSoftwareVersion/w_softwareVersion",
    "Trusted Certificates": "/cgi-bin/cm/secTrustedCertificates/w_trustedCertificates",
    "Server/Application Certificates": "/cgi-bin/cm/secServerCertificates/w_serverCertificates"
}

sat_commands = [
    "status media-processor all",
    "list media-gateway",
    "list survivable-processor",
    "status aesvcs cti-link",
    "status processor-channels 3",
    "status processor-channels 5",
    "list measurements outage-trunk last-hour",
    "status aesvcs interface",
    "status aesvcs link",
    "status cdr-link"
]

linux_commands = [
    'statapp',
    'date',
    'uptime',
    'df -h',
    'df -k',
    'cat /etc/hosts'
]

def clean_output(output):
    output = re.sub(r'\x1B[@-_][0-?]*[ -/]*[@-~]', '', output)
    output = re.sub(r'^Command:.*$', '', output, flags=re.MULTILINE)
    output = re.sub(r'\n{2,}', '\n\n', output)
    output = ''.join(filter(lambda x: x in string.printable, output))
    return output.strip()

def auth_handler(title, instructions, prompt_list):
    return [SAT_PASSWORD if 'Password' in p[0] else '' for p in prompt_list]

# Excel setup
if not os.path.exists(excel_path):
    wb = openpyxl.Workbook()
    wb.save(excel_path)

wb = openpyxl.load_workbook(excel_path)
for sheet in wb.sheetnames:
    del wb[sheet]

wb.create_sheet("Web GUI")
wb.create_sheet("SAT")
wb.create_sheet("Linux")

# Web GUI Section
try:
    chrome_options = Options()
    chrome_options.add_experimental_option("detach", True)
    chrome_options.add_argument("--start-maximized")
    chrome_options.add_argument("--ignore-certificate-errors")

    service = Service(CHROME_DRIVER_PATH)
    driver = webdriver.Chrome(service=service, options=chrome_options)
    wait = WebDriverWait(driver, 15)

    driver.get(LOGIN_URL)
    wait.until(EC.presence_of_element_located((By.NAME, "userName"))).send_keys(WEB_USERNAME)
    driver.find_element(By.NAME, "logonButton").click()
    wait.until(EC.presence_of_element_located((By.NAME, "pa55word"))).send_keys(WEB_PASSWORD)
    driver.find_element(By.NAME, "logonButton").click()
    wait.until(EC.element_to_be_clickable((By.NAME, "motdContinue"))).click()

    admin_menu = wait.until(EC.presence_of_element_located((By.XPATH, "//a[text()='Administration']")))
    ActionChains(driver).move_to_element(admin_menu).pause(1).click().perform()

    server_maintenance = wait.until(EC.element_to_be_clickable((
        By.XPATH,
        "//a[contains(@href, \"setBreadCrumb('/cgi-bin/cm/server/w_server'\")]"
    )))
    server_maintenance.click()

    ws = wb["Web GUI"]
    if ws.max_row == 1:
        ws.append(["Timestamp", "Section Heading", "Content"])

    for label, url in web_pages.items():
        full_url = f"https://172.16.70.39{url}"
        driver.get(full_url)
        wait.until(EC.presence_of_element_located((By.ID, "mainWorkArea")))
        heading = driver.find_element(By.CLASS_NAME, "heading1Plain").text
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        if label in ["Trusted Certificates", "Server/Application Certificates"]:
            rows = driver.find_elements(By.XPATH, "//table//tr[position()>1]")
            formatted = []
            for row in rows:
                cells = row.find_elements(By.TAG_NAME, "td")
                if len(cells) >= 5:
                    cert_name = cells[1].text.strip()
                    expiry_date = cells[4].text.strip()
                    if cert_name and expiry_date:
                        formatted.append(f"{cert_name} certificate is expiring on {expiry_date}")
            content = "\n".join(formatted)
        else:
            content = driver.find_element(By.ID, "mainWorkArea").text.strip()

        ws.append([timestamp, heading, content])
        ws[f"C{ws.max_row}"].alignment = Alignment(wrap_text=True)

    driver.quit()
    print("✅ GUI data fetched")
except Exception as e:
    print(f"❌ Web scraping error: {e}")
    driver.quit()

# SAT SSH
try:
    transport = paramiko.Transport((SAT_HOST, SAT_PORT))
    transport.connect()
    transport.auth_interactive(SAT_USERNAME, auth_handler)
    channel = transport.open_session()
    channel.get_pty()
    channel.invoke_shell()
    time.sleep(2)
    if channel.recv_ready():
        channel.recv(9999)
    channel.send("sat\n")
    time.sleep(3)
    if "Terminal Type" in channel.recv(4096).decode(errors='ignore'):
        channel.send("VT220\n")
        time.sleep(1)
        if channel.recv_ready():
            channel.recv(9999)

    ws = wb["SAT"]
    if ws.max_row == 1:
        ws.append(["Timestamp", "SAT Command", "Output"])

    for command in sat_commands:
        print(f"→ Running SAT: {command}")
        channel.send(command + "\n")
        time.sleep(2)
        result = ""
        while channel.recv_ready():
            result += channel.recv(4096).decode(errors='ignore')
        cleaned = clean_output(result)
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.append([timestamp, command, cleaned])
    transport.close()
    print("✅ SAT commands completed")
except Exception as e:
    print(f"❌ SAT SSH Error: {e}")

# Linux SSH
try:
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    client.connect(SSH_HOST, SSH_PORT, SSH_USERNAME, SSH_PASSWORD)

    ws = wb["Linux"]
    if ws.max_row == 1:
        ws.append(["Timestamp", "Command", "Output", "Error"])

    for command in linux_commands:
        print(f"→ Running SSH: {command}")
        stdin, stdout, stderr = client.exec_command(command)
        output = stdout.read().decode().strip()
        error = stderr.read().decode().strip()
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.append([timestamp, command, output, error])
    client.close()
    print("✅ SSH commands completed")
except Exception as e:
    print(f"❌ SSH Error: {e}")

wb.save(excel_path)
print(f"\n✅ All data saved to '{excel_path}'")


# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:




